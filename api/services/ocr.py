import asyncio
import base64
import json
import logging
import subprocess
import tempfile
from collections.abc import Awaitable, Callable
from pathlib import Path

import asyncpg
import httpx
from botocore.exceptions import BotoCoreError
from config import settings
from infra.db.derived_documents import replace_derived_content
from services.chunker import chunk_pages, chunk_text
from services.extracted_assets import ExtractedAsset, build_pdf_image_assets
from services.pdf_extract import extract_pdf
from services.s3 import S3Service

logger = logging.getLogger(__name__)

MISTRAL_OCR_URL = "https://api.mistral.ai/v1/ocr"
MAX_RETRIES = 3
RETRY_BACKOFF = [2, 5, 10]

OFFICE_TYPES = {"pptx", "ppt", "docx", "doc"}
IMAGE_TYPES = {"png", "jpg", "jpeg", "webp", "gif"}
OCR_TYPES = {"pdf"} | OFFICE_TYPES | IMAGE_TYPES

BeforeWrite = Callable[[asyncpg.Connection], Awaitable[None]]
BeforeArtifactWrite = Callable[[], Awaitable[None]]
ArtifactObject = tuple[str, bytes, str]
_ARTIFACT_POINTER_KEYS = frozenset({"converted_s3_key", "ocr_s3_key", "tagged_s3_key"})


class ExtractionError(RuntimeError):
    """Stable, sanitized extraction failure for durable job classification."""

    def __init__(self, error_code: str, error_message: str) -> None:
        self.error_code = error_code
        self.error_message = error_message
        super().__init__(error_message)


class RetryableExtractionError(ExtractionError):
    """A transient storage, network, converter, or database extraction failure."""


class TerminalExtractionError(ExtractionError):
    """An invalid document or quota condition that retrying cannot fix."""


class OCRService:
    def __init__(self, s3: S3Service, pool: asyncpg.Pool):
        self._s3 = s3
        self._pool = pool
        self._semaphore = asyncio.Semaphore(3)

    async def process_document(self, document_id: str, user_id: str):
        """Legacy rollback path that owns its own failure transition."""
        async with self._semaphore:
            try:
                await self._do_process(document_id, user_id)
            except Exception as exc:  # noqa: BLE001 - legacy path records a stable failure.
                logger.error(
                    "Document processing failed document_id=%s error_type=%s",
                    document_id,
                    type(exc).__name__,
                )
                try:
                    await self._pool.execute(
                        "UPDATE documents SET status = 'failed', error_message = $2, updated_at = now() WHERE id = $1",
                        document_id,
                        "Document processing failed.",
                    )
                except Exception as update_exc:  # noqa: BLE001 - keep legacy task contained.
                    logger.error(
                        "Document failure status update failed document_id=%s error_type=%s",
                        document_id,
                        type(update_exc).__name__,
                    )

    async def extract_document(
        self,
        document_id: str,
        user_id: str,
        *,
        before_write: BeforeWrite,
        before_artifact_write: BeforeArtifactWrite,
        artifact_namespace: str | None = None,
    ) -> int:
        """Run extraction transparently for a durable worker-owned lease."""
        async with self._semaphore:
            try:
                return await self._do_process(
                    document_id,
                    user_id,
                    before_write=before_write,
                    before_artifact_write=before_artifact_write,
                    artifact_namespace=artifact_namespace,
                    set_processing=False,
                )
            except (TerminalExtractionError, RetryableExtractionError):
                raise
            except LookupError:
                raise TerminalExtractionError(
                    "document_not_found",
                    "Document was not found.",
                ) from None
            except (asyncpg.PostgresError, BotoCoreError, httpx.TransportError, OSError, TimeoutError):
                raise RetryableExtractionError(
                    "extraction_transient",
                    "Document extraction will be retried.",
                ) from None

    async def _check_global_limits(self, document_id: str):
        if not settings.GLOBAL_OCR_ENABLED:
            raise TerminalExtractionError(
                "extraction_disabled",
                "Document extraction is disabled.",
            )

        total_pages = await self._pool.fetchval("SELECT COALESCE(SUM(page_count), 0) FROM documents WHERE NOT archived")
        if total_pages >= settings.GLOBAL_MAX_PAGES:
            raise TerminalExtractionError(
                "quota_exceeded",
                "Document extraction quota was exceeded.",
            )

    async def _check_user_page_limit(
        self,
        user_id: str,
        new_pages: int,
        conn: asyncpg.Connection | None = None,
        *,
        after_lock: BeforeWrite | None = None,
    ):
        """Quota check — uses an advisory lock when given a transaction connection so concurrent jobs serialize."""
        executor = conn or self._pool
        if after_lock is not None and conn is None:
            raise ValueError("after_lock requires a transaction connection")
        if conn is not None:
            # pg_advisory_xact_lock serializes concurrent OCR jobs for the same
            # user inside this transaction; releases on commit/rollback.
            await conn.execute(
                "SELECT pg_advisory_xact_lock(hashtext($1::text))",
                user_id,
            )
        if after_lock is not None:
            await after_lock(conn)
        row = await executor.fetchrow(
            "SELECT u.page_limit, "
            "COALESCE((SELECT SUM(page_count) FROM documents WHERE user_id = $1 AND NOT archived), 0)::bigint AS used "
            "FROM users u WHERE u.id = $1",
            user_id,
        )
        if not row:
            return
        limit = row["page_limit"]
        used = row["used"] or 0
        if used + new_pages > limit:
            raise TerminalExtractionError(
                "quota_exceeded",
                "Document extraction quota was exceeded.",
            )

    async def _do_process(
        self,
        document_id: str,
        user_id: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
        set_processing: bool = True,
    ) -> int:
        await self._check_global_limits(document_id)
        if set_processing:
            await self._set_status(document_id, "processing")

        doc = await self._pool.fetchrow(
            "SELECT filename, file_type, path, knowledge_base_id::text as kb_id "
            "FROM documents WHERE id = $1 AND user_id = $2",
            document_id,
            user_id,
        )
        if not doc:
            raise TerminalExtractionError("document_not_found", "Document was not found.")

        ext = doc["filename"].rsplit(".", 1)[-1].lower() if "." in doc["filename"] else doc["file_type"]
        kb_id = doc["kb_id"]
        s3_source_key = f"{user_id}/{document_id}/source.{ext}"

        write_kwargs = {"before_write": before_write} if before_write is not None else {}
        durable_kwargs = dict(write_kwargs)
        if before_artifact_write is not None:
            durable_kwargs["before_artifact_write"] = before_artifact_write
        if artifact_namespace is not None:
            durable_kwargs["artifact_namespace"] = artifact_namespace
        if ext in OFFICE_TYPES:
            return await self._process_office(document_id, user_id, kb_id, s3_source_key, ext, **durable_kwargs)
        if ext in IMAGE_TYPES:
            return await self._process_image(document_id, user_id, kb_id, s3_source_key, ext, **write_kwargs)
        if ext == "pdf":
            return await self._process_pdf(document_id, user_id, kb_id, s3_source_key, **durable_kwargs)
        if ext in ("html", "htm"):
            return await self._process_html(document_id, user_id, kb_id, s3_source_key, **durable_kwargs)
        if ext in ("xlsx", "xls", "csv"):
            return await self._process_spreadsheet(document_id, user_id, kb_id, s3_source_key, ext, **write_kwargs)
        raise TerminalExtractionError(
            "unsupported_document_type",
            "Document type is not supported.",
        )

    # ── PDF extraction ────────────────────────────────────────────────────

    async def _process_pdf(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        if settings.PDF_BACKEND == "mistral":
            if not settings.MISTRAL_API_KEY:
                raise TerminalExtractionError("extraction_configuration", "Document extraction is unavailable.")
            # Cheap pre-check: refuse if user is already over quota so we don't burn the Mistral call.
            await self._check_user_page_limit(user_id, 1)
            presigned_url = await self._s3.generate_presigned_get(s3_source_key)
            ocr_result = await self._call_mistral_ocr(presigned_url, "document_url")
            return await self._store_ocr_result(
                document_id,
                user_id,
                kb_id,
                ocr_result,
                **({"before_write": before_write} if before_write is not None else {}),
                **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
                **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
            )
        if settings.CONVERTER_URL:
            presigned_url = await self._s3.generate_presigned_get(s3_source_key)
            pages = await self._call_converter_extract(presigned_url, "pdf")
            return await self._store_extracted_pages(
                document_id,
                user_id,
                kb_id,
                pages,
                "opendataloader",
                **({"before_write": before_write} if before_write is not None else {}),
                **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
                **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
            )
        return await self._process_opendataloader(
            document_id,
            user_id,
            kb_id,
            s3_source_key,
            **({"before_write": before_write} if before_write is not None else {}),
            **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
            **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
        )

    async def _process_office(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        ext: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        """Process Office files. Routes through converter or falls back to local LibreOffice."""
        if settings.PDF_BACKEND == "mistral":
            attempt_pdf_key = (
                f"{user_id}/{document_id}/derived/{artifact_namespace}/converted.pdf"
                if artifact_namespace is not None
                else f"{user_id}/{document_id}/converted.pdf"
            )
            try:
                if before_artifact_write is not None:
                    await before_artifact_write()
                pdf_key = await self._convert_to_pdf_s3(
                    document_id,
                    user_id,
                    s3_source_key,
                    ext,
                    artifact_namespace=artifact_namespace,
                )
                if before_artifact_write is not None:
                    await before_artifact_write()
            except BaseException:  # noqa: BLE001 - converter may PUT before returning an error or cancellation.
                if artifact_namespace is not None:
                    await self._delete_artifact_keys([attempt_pdf_key])
                raise
            converted_metadata = {"converted_s3_key": pdf_key if artifact_namespace is not None else None}
            try:
                if not settings.MISTRAL_API_KEY:
                    raise TerminalExtractionError("extraction_configuration", "Document extraction is unavailable.")
                await self._check_user_page_limit(user_id, 1)
                presigned_url = await self._s3.generate_presigned_get(pdf_key)
                ocr_result = await self._call_mistral_ocr(presigned_url, "document_url")
                return await self._store_ocr_result(
                    document_id,
                    user_id,
                    kb_id,
                    ocr_result,
                    **({"before_write": before_write} if before_write is not None else {}),
                    **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
                    **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
                    metadata_patch_extra=converted_metadata,
                )
            except BaseException:  # noqa: BLE001 - compensate cancellation and uncertain final publication.
                if artifact_namespace is not None:
                    published = await self._artifacts_were_published(
                        document_id,
                        user_id,
                        converted_metadata,
                        [],
                    )
                    if published is False:
                        await self._delete_artifact_keys([pdf_key])
                raise
        if settings.CONVERTER_URL:
            presigned_url = await self._s3.generate_presigned_get(s3_source_key)
            pages = await self._call_converter_extract(presigned_url, ext)
            return await self._store_extracted_pages(
                document_id,
                user_id,
                kb_id,
                pages,
                "opendataloader",
                **({"before_write": before_write} if before_write is not None else {}),
                **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
                **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
            )
        return await self._process_office_local(
            document_id,
            user_id,
            kb_id,
            s3_source_key,
            ext,
            **({"before_write": before_write} if before_write is not None else {}),
            **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
            **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
        )

    # ── Converter integration (hosted mode) ───────────────────────────────

    async def _call_converter_extract(self, source_url: str, ext: str) -> list[tuple[int, str]]:
        """Call the converter /extract endpoint. Returns list of (page_num, markdown).

        Sends a request_id for source binding. If the converter echoes it back,
        we verify the match; if the converter doesn't support it, we log a warning
        but still accept the response (forward-compatible).
        """
        import uuid as _uuid

        request_id = str(_uuid.uuid4())

        headers = {}
        if settings.CONVERTER_SECRET:
            headers["Authorization"] = f"Bearer {settings.CONVERTER_SECRET}"

        async with httpx.AsyncClient(timeout=httpx.Timeout(300.0, connect=10.0)) as client:
            resp = await client.post(
                f"{settings.CONVERTER_URL}/extract",
                json={"source_url": source_url, "source_ext": ext, "request_id": request_id},
                headers=headers,
            )
            resp.raise_for_status()
            data = resp.json()

        # Source binding: verify request_id echo if the converter supports it
        echoed_id = data.get("request_id")
        if echoed_id is not None and echoed_id != request_id:
            raise TerminalExtractionError(
                "invalid_document",
                "Document extraction response was invalid.",
            )
        if echoed_id is None:
            logger.warning("Converter did not echo request_id — source binding not verified")

        pages = data.get("pages", [])
        if not pages:
            raise TerminalExtractionError("invalid_document", "Document content is invalid.")

        return [(p["page"], p["content"]) for p in pages]

    # ── OpenDataLoader local extraction ───────────────────────────────────

    async def _process_opendataloader(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        """Extract PDF via opendataloader-pdf (local mode or hosted fallback)."""
        with tempfile.TemporaryDirectory() as tmpdir:
            pdf_path = Path(tmpdir) / "source.pdf"
            await self._s3.download_to_file(s3_source_key, str(pdf_path))
            pages_with_images = await asyncio.to_thread(extract_pdf, str(pdf_path))

        assets, page_elements = await self._build_pdf_assets(document_id, pages_with_images)

        page_contents = [(num, md) for num, md, _ in pages_with_images]
        version = await self._store_extracted_pages(
            document_id,
            user_id,
            kb_id,
            page_contents,
            "opendataloader",
            page_elements=page_elements,
            assets=assets,
            **({"before_write": before_write} if before_write is not None else {}),
            **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
            **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
        )
        if artifact_namespace is None:
            await self._upload_assets(user_id, assets)
        return version

    # ── Office local fallback (no converter) ──────────────────────────────

    async def _convert_to_pdf_s3(
        self,
        document_id: str,
        user_id: str,
        s3_source_key: str,
        ext: str,
        *,
        artifact_namespace: str | None = None,
    ) -> str:
        """Convert Office file to PDF and upload to S3. Returns S3 key of the PDF."""
        pdf_key = f"{user_id}/{document_id}/converted.pdf"
        if artifact_namespace is not None:
            pdf_key = f"{user_id}/{document_id}/derived/{artifact_namespace}/converted.pdf"

        if settings.CONVERTER_URL:
            # Legacy path — only used for Mistral backend with converter
            import uuid as _uuid

            request_id = str(_uuid.uuid4())
            source_url = await self._s3.generate_presigned_get(s3_source_key)
            result_url = await self._s3.generate_presigned_put(pdf_key)
            headers = {}
            if settings.CONVERTER_SECRET:
                headers["Authorization"] = f"Bearer {settings.CONVERTER_SECRET}"
            async with httpx.AsyncClient(timeout=httpx.Timeout(180.0, connect=10.0)) as client:
                resp = await client.post(
                    f"{settings.CONVERTER_URL}/convert",
                    json={
                        "source_url": source_url,
                        "result_url": result_url,
                        "source_ext": ext,
                        "request_id": request_id,
                    },
                    headers=headers,
                )
                resp.raise_for_status()
                data = resp.json() if resp.headers.get("content-type", "").startswith("application/json") else {}
                echoed_id = data.get("request_id") if isinstance(data, dict) else None
                if echoed_id is not None and echoed_id != request_id:
                    raise ValueError(f"Converter response binding mismatch: sent {request_id}, got {echoed_id}")
        else:
            with tempfile.TemporaryDirectory() as tmpdir:
                source_path = Path(tmpdir) / f"source.{ext}"
                await self._s3.download_to_file(s3_source_key, str(source_path))
                result = await asyncio.to_thread(
                    subprocess.run,
                    [
                        "libreoffice",
                        "--headless",
                        "--norestore",
                        f"-env:UserInstallation=file://{tmpdir}/lo-profile",  # 并发转换各用独立配置,防互踩锁
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        tmpdir,
                        str(source_path),
                    ],
                    capture_output=True,
                    timeout=120,
                )
                if result.returncode != 0:
                    raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.decode()[:300]}")
                pdf_path = Path(tmpdir) / "source.pdf"
                if not pdf_path.exists():
                    raise RuntimeError("LibreOffice did not produce a PDF")
                await self._s3.upload_file(pdf_key, str(pdf_path), "application/pdf")

        return pdf_key

    async def _process_office_local(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        ext: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        """Convert Office file to PDF locally, then extract with opendataloader."""
        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / f"source.{ext}"
            await self._s3.download_to_file(s3_source_key, str(source_path))

            result = await asyncio.to_thread(
                subprocess.run,
                [
                    "libreoffice",
                    "--headless",
                    "--norestore",
                    f"-env:UserInstallation=file://{tmpdir}/lo-profile",  # 并发转换各用独立配置,防互踩锁
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    tmpdir,
                    str(source_path),
                ],
                capture_output=True,
                timeout=120,
            )
            if result.returncode != 0:
                raise RuntimeError(f"LibreOffice conversion failed: {result.stderr.decode()[:300]}")

            pdf_path = Path(tmpdir) / "source.pdf"
            if not pdf_path.exists():
                raise RuntimeError("LibreOffice did not produce a PDF")

            pages_with_images = await asyncio.to_thread(extract_pdf, str(pdf_path))

        assets, page_elements = await self._build_pdf_assets(document_id, pages_with_images)

        page_contents = [(num, md) for num, md, _ in pages_with_images]
        version = await self._store_extracted_pages(
            document_id,
            user_id,
            kb_id,
            page_contents,
            "libreoffice+opendataloader",
            page_elements=page_elements,
            assets=assets,
            **({"before_write": before_write} if before_write is not None else {}),
            **({"before_artifact_write": before_artifact_write} if before_artifact_write is not None else {}),
            **({"artifact_namespace": artifact_namespace} if artifact_namespace is not None else {}),
        )
        if artifact_namespace is None:
            await self._upload_assets(user_id, assets)
        return version

    # ── Shared page storage ───────────────────────────────────────────────

    async def _build_pdf_assets(
        self,
        document_id: str,
        pages_with_images: list[tuple[int, str, list[dict]]],
    ) -> tuple[list[ExtractedAsset], dict[int, dict]]:
        doc = await self._pool.fetchrow(
            "SELECT filename, path FROM documents WHERE id = $1",
            document_id,
        )
        if not doc:
            return [], {}
        return build_pdf_image_assets(
            document_id,
            doc["filename"],
            doc["path"],
            pages_with_images,
        )

    async def _upload_assets(self, user_id: str, assets: list[ExtractedAsset]) -> None:
        for asset in assets:
            await self._s3.upload_bytes(
                f"{user_id}/{asset.document_id}/source.{asset.file_type}",
                asset.data,
                asset.content_type,
            )

    async def _commit_derived_content(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        *,
        pages: list[tuple[int, str]],
        chunks,
        parser: str,
        content: str | None = None,
        page_elements: dict[int, dict] | None = None,
        assets: list[ExtractedAsset] | None = None,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
        artifact_objects: list[ArtifactObject] | None = None,
        metadata_patch_extra: dict | None = None,
    ) -> int:
        pending_artifacts = list(artifact_objects or [])
        if artifact_namespace is not None and assets:
            pending_artifacts.extend(
                (
                    f"{user_id}/{asset.document_id}/source.{asset.file_type}",
                    asset.data,
                    asset.content_type,
                )
                for asset in assets
            )
        uploaded_keys: list[str] = []

        async def check_page_limit(conn):
            await self._check_user_page_limit(
                user_id,
                len(pages),
                conn=conn,
                after_lock=before_write,
            )

        metadata_patch = dict(metadata_patch_extra or {})
        if assets is not None:
            metadata_patch["assets"] = [asset.metadata() for asset in assets]

        try:
            for key, data, content_type in pending_artifacts:
                if before_artifact_write is not None:
                    await before_artifact_write()
                uploaded_keys.append(key)
                await self._s3.upload_bytes(key, data, content_type)
                if before_artifact_write is not None:
                    await before_artifact_write()

            current_artifact_keys = {key for key, _, _ in pending_artifacts}
            current_artifact_keys.update(
                value
                for key, value in (metadata_patch_extra or {}).items()
                if key in _ARTIFACT_POINTER_KEYS and isinstance(value, str) and value
            )

            async def cleanup_replaced_artifacts(keys: list[str]) -> None:
                await self._delete_artifact_keys([key for key in keys if key not in current_artifact_keys])

            return await replace_derived_content(
                self._pool,
                document_id=document_id,
                user_id=user_id,
                knowledge_base_id=kb_id,
                pages=pages,
                chunks=chunks,
                parser=parser,
                content=content,
                page_elements=page_elements,
                metadata_patch=metadata_patch,
                assets=assets,
                before_write=check_page_limit,
                after_commit=cleanup_replaced_artifacts,
            )
        except BaseException:
            published = await self._artifacts_were_published(
                document_id,
                user_id,
                metadata_patch_extra or {},
                assets or [],
            )
            if published is False:
                await self._delete_artifact_keys(uploaded_keys)
            raise

    async def _artifacts_were_published(
        self,
        document_id: str,
        user_id: str,
        metadata_values: dict,
        assets: list[ExtractedAsset],
    ) -> bool | None:
        try:
            row = await self._pool.fetchrow(
                "SELECT status::text, metadata FROM documents WHERE id = $1 AND user_id = $2",
                document_id,
                user_id,
            )
            if row is None or row["status"] != "ready":
                return False
            metadata = row["metadata"]
            if isinstance(metadata, str):
                metadata = json.loads(metadata)
            if any((metadata or {}).get(key) != value for key, value in metadata_values.items()):
                return False
            if assets:
                count = await self._pool.fetchval(
                    "SELECT count(*) FROM documents WHERE id = ANY($1::uuid[]) AND user_id = $2",
                    [asset.document_id for asset in assets],
                    user_id,
                )
                if count != len(assets):
                    return False
            return True
        except Exception:  # noqa: BLE001 - an uncertain commit must preserve artifacts.
            return None

    async def _delete_artifact_keys(self, keys: list[str]) -> None:
        for key in keys:
            try:
                delete_key = getattr(self._s3, "delete_key", None)
                if delete_key is not None:
                    await delete_key(key)
                else:
                    await self._s3.delete_prefix(key)
            except Exception as exc:  # noqa: BLE001 - cleanup is best effort after confirmed rollback.
                logger.error("Derived artifact cleanup failed error_type=%s", type(exc).__name__)

    async def _store_extracted_pages(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        page_contents: list[tuple[int, str]],
        parser: str,
        page_elements: dict[int, dict] | None = None,
        assets: list[ExtractedAsset] | None = None,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        """Store pages/chunks and update document status."""
        num_pages = len(page_contents)

        if num_pages > settings.QUOTA_MAX_PAGES_PER_DOC:
            raise TerminalExtractionError(
                "quota_exceeded",
                "Document extraction quota was exceeded.",
            )

        full_content = "\n\n---\n\n".join(md for _, md in page_contents)
        chunks = chunk_pages(page_contents)

        version = await self._commit_derived_content(
            document_id,
            user_id,
            kb_id,
            pages=page_contents,
            chunks=chunks,
            parser=parser,
            content=full_content,
            page_elements=page_elements,
            assets=assets or [],
            before_write=before_write,
            before_artifact_write=before_artifact_write,
            artifact_namespace=artifact_namespace,
        )
        logger.info("Extracted (%s): doc=%s pages=%d chunks=%d", parser, document_id[:8], num_pages, len(chunks))
        return version

    # ── Image processing ──────────────────────────────────────────────────

    async def _process_image(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        ext: str,
        *,
        before_write: BeforeWrite | None = None,
    ) -> int:
        """Images are stored as-is. No OCR. The MCP read tool returns them natively."""
        conn = await self._pool.acquire()
        try:
            async with conn.transaction():
                await self._check_user_page_limit(user_id, 1, conn=conn, after_lock=before_write)
                version = await conn.fetchval(
                    "UPDATE documents SET status = 'ready', page_count = 1, parser = 'native', "
                    "version = version + 1, error_message = NULL, updated_at = now() "
                    "WHERE id = $1 AND user_id = $2 AND knowledge_base_id = $3 "
                    "AND NOT archived AND source_kind = 'source' RETURNING version",
                    document_id,
                    user_id,
                    kb_id,
                )
        finally:
            await self._pool.release(conn)
        logger.info("Image stored: doc=%s", document_id[:8])
        if version is None:
            raise TerminalExtractionError("document_not_found", "Document was not found.")
        return version

    # ── HTML processing ───────────────────────────────────────────────────

    async def _process_html(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
    ) -> int:
        """Parse HTML with webmd parser, store markdown + tagged HTML."""
        from html_parser import Parser

        html_bytes = await self._s3.download_bytes(s3_source_key)
        raw_html = html_bytes.decode("utf-8", errors="replace")

        parser = Parser(raw_html, content_only=True)
        result = parser.parse()

        await parser.embed_images()
        tagged_html = parser.html(sanitize=True)

        tagged_bytes = tagged_html.encode("utf-8")
        tagged_key = f"{user_id}/{document_id}/tagged.html"
        artifact_objects = None
        metadata_patch_extra = None
        if artifact_namespace is None:
            await self._s3.upload_bytes(tagged_key, tagged_bytes, "text/html")
            metadata_patch_extra = {"tagged_s3_key": None}
        else:
            tagged_key = f"{user_id}/{document_id}/derived/{artifact_namespace}/tagged.html"
            artifact_objects = [(tagged_key, tagged_bytes, "text/html")]
            metadata_patch_extra = {"tagged_s3_key": tagged_key}

        markdown_content = result.content
        chunks = chunk_text(markdown_content)

        version = await self._commit_derived_content(
            document_id,
            user_id,
            kb_id,
            pages=[(1, markdown_content)],
            chunks=chunks,
            parser="webmd",
            content=markdown_content,
            before_write=before_write,
            before_artifact_write=before_artifact_write,
            artifact_namespace=artifact_namespace,
            artifact_objects=artifact_objects,
            metadata_patch_extra=metadata_patch_extra,
        )
        logger.info("HTML processed: doc=%s chunks=%d", document_id[:8], len(chunks))
        return version

    # ── Spreadsheet processing ────────────────────────────────────────────

    async def _process_spreadsheet(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        s3_source_key: str,
        ext: str,
        *,
        before_write: BeforeWrite | None = None,
    ) -> int:
        """Download spreadsheet, store each sheet as a document_page."""
        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / f"source.{ext}"
            await self._s3.download_to_file(s3_source_key, str(source_path))

            sheets = await asyncio.to_thread(self._parse_sheets, str(source_path), ext)

            content_parts = [f"## {name}\n\n{md}" for name, md in sheets]
            full_content = "\n\n---\n\n".join(content_parts)
            page_contents = [(i + 1, md) for i, (_, md) in enumerate(sheets)]
            chunks = chunk_pages(page_contents)

            version = await self._commit_derived_content(
                document_id,
                user_id,
                kb_id,
                pages=page_contents,
                chunks=chunks,
                parser="openpyxl",
                content=full_content,
                page_elements={index: {"sheet_name": name} for index, (name, _) in enumerate(sheets, 1)},
                before_write=before_write,
            )
            logger.info("Spreadsheet processed: doc=%s sheets=%d chunks=%d", document_id[:8], len(sheets), len(chunks))
            return version

    @staticmethod
    def _rows_to_markdown(rows: list[list[str]], max_rows: int = 100) -> str:
        if not rows:
            return "(empty)"
        header = "| " + " | ".join(rows[0]) + " |"
        sep = "| " + " | ".join("---" for _ in rows[0]) + " |"
        data = rows[1 : max_rows + 1]
        body = "\n".join("| " + " | ".join(r) + " |" for r in data)
        truncated = f"\n\n*({len(rows) - 1 - max_rows} more rows truncated)*" if len(rows) - 1 > max_rows else ""
        return f"{header}\n{sep}\n{body}{truncated}"

    @staticmethod
    def _parse_sheets(path: str, ext: str) -> list[tuple[str, str]]:
        """Returns list of (sheet_name, markdown_table) tuples."""
        import csv

        if ext == "csv":
            with open(path, newline="", encoding="utf-8", errors="replace") as f:
                rows = [[c for c in row] for row in csv.reader(f)]
            return [("Sheet1", OCRService._rows_to_markdown(rows))]

        if ext == "xls":  # 老式二进制格式,openpyxl 不支持
            try:
                import xlrd
            except ImportError:
                return [("Error", "(xlrd not installed)")]
            book = xlrd.open_workbook(path)
            return [
                (
                    sheet.name,
                    OCRService._rows_to_markdown(
                        [["" if c.value is None else str(c.value) for c in sheet.row(r)] for r in range(sheet.nrows)]
                    ),
                )
                for sheet in book.sheets()
                if sheet.nrows
            ]

        try:
            import openpyxl
        except ImportError:
            return [("Error", "(openpyxl not installed)")]

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for name in wb.sheetnames:
            ws = wb[name]
            rows = [[str(cell) if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]
            if not rows:
                continue
            sheets.append((name, OCRService._rows_to_markdown(rows)))
        wb.close()
        return sheets

    # ── Mistral OCR ───────────────────────────────────────────────────────

    @staticmethod
    def _stored_ocr_page_elements(pages: list[dict], page_elements: dict[int, dict]) -> dict[int, dict]:
        stored_elements: dict[int, dict] = {}
        for page in pages:
            page_index = page.get("index", 0) + 1
            elements = {}
            page_assets = page_elements.get(page_index, {}).get("images")
            if page_assets:
                elements["images"] = page_assets
            if page.get("dimensions"):
                elements["dimensions"] = page["dimensions"]
            if page.get("tables"):
                elements["tables"] = page["tables"]
            if elements:
                stored_elements[page_index] = elements
        return stored_elements

    async def _store_ocr_result(
        self,
        document_id: str,
        user_id: str,
        kb_id: str,
        ocr_result: dict,
        *,
        before_write: BeforeWrite | None = None,
        before_artifact_write: BeforeArtifactWrite | None = None,
        artifact_namespace: str | None = None,
        metadata_patch_extra: dict | None = None,
    ) -> int:
        ocr_json_bytes = json.dumps(ocr_result).encode()
        ocr_key = f"{user_id}/{document_id}/ocr.json"
        artifact_objects = None
        metadata_patch = dict(metadata_patch_extra or {})
        if artifact_namespace is None:
            await self._s3.upload_bytes(ocr_key, ocr_json_bytes, "application/json")
            metadata_patch["ocr_s3_key"] = None
        else:
            ocr_key = f"{user_id}/{document_id}/derived/{artifact_namespace}/ocr.json"
            artifact_objects = [(ocr_key, ocr_json_bytes, "application/json")]
            metadata_patch["ocr_s3_key"] = ocr_key

        pages = ocr_result.get("pages", [])

        if len(pages) > settings.QUOTA_MAX_PAGES_PER_DOC:
            raise TerminalExtractionError(
                "quota_exceeded",
                "Document extraction quota was exceeded.",
            )

        pages_with_images = []
        for page in pages:
            page_num = page.get("index", 0) + 1
            extracted_images = []
            for img in page.get("images", []):
                img_id = img.get("id")
                img_b64 = img.get("image_base64")
                if not img_id or not img_b64:
                    continue
                if img_b64.startswith("data:"):
                    img_b64 = img_b64.split(",", 1)[1]
                img_bytes = base64.b64decode(img_b64)
                extracted_images.append(
                    {
                        "id": img_id,
                        "bytes": img_bytes,
                        "format": "jpg",
                    }
                )
            pages_with_images.append((page_num, page.get("markdown", ""), extracted_images))

        assets, page_elements = await self._build_pdf_assets(document_id, pages_with_images)

        page_count = len(pages)
        content_parts = [page.get("markdown", "") for page in pages]
        full_content = "\n\n---\n\n".join(content_parts)
        page_contents = [(page.get("index", 0) + 1, page.get("markdown", "")) for page in pages]
        chunks = chunk_pages(page_contents)

        stored_elements = self._stored_ocr_page_elements(pages, page_elements)

        version = await self._commit_derived_content(
            document_id,
            user_id,
            kb_id,
            pages=page_contents,
            chunks=chunks,
            parser="mistral",
            content=full_content,
            page_elements=stored_elements,
            assets=assets,
            before_write=before_write,
            before_artifact_write=before_artifact_write,
            artifact_namespace=artifact_namespace,
            artifact_objects=artifact_objects,
            metadata_patch_extra=metadata_patch,
        )

        if artifact_namespace is None:
            await self._upload_assets(user_id, assets)
        logger.info("OCR complete: doc=%s pages=%d chunks=%d", document_id[:8], page_count, len(chunks))
        return version

    async def _call_mistral_ocr(self, url: str, url_type: str = "document_url") -> dict:
        last_error = None
        for attempt in range(MAX_RETRIES):
            try:
                async with httpx.AsyncClient(timeout=httpx.Timeout(300.0, connect=10.0)) as client:
                    resp = await client.post(
                        MISTRAL_OCR_URL,
                        headers={
                            "Authorization": f"Bearer {settings.MISTRAL_API_KEY}",
                            "Content-Type": "application/json",
                        },
                        json={
                            "model": "mistral-ocr-latest",
                            "document": {
                                "type": url_type,
                                url_type: url,
                            },
                            "include_image_base64": True,
                            "table_format": "markdown",
                        },
                    )
                    resp.raise_for_status()
                    return resp.json()
            except (httpx.HTTPStatusError, httpx.TimeoutException) as e:
                last_error = e
                if attempt < MAX_RETRIES - 1:
                    wait = RETRY_BACKOFF[attempt]
                    logger.warning(
                        "Mistral OCR attempt failed attempt=%d error_type=%s retry_in_seconds=%d",
                        attempt + 1,
                        type(e).__name__,
                        wait,
                    )
                    await asyncio.sleep(wait)
        raise last_error or RuntimeError("Mistral OCR failed after retries")

    async def _set_status(self, document_id: str, status: str):
        await self._pool.execute(
            "UPDATE documents SET status = $2, updated_at = now() WHERE id = $1",
            document_id,
            status,
        )
